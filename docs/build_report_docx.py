from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "NOI_DUNG_BAO_CAO_HEALWAY.md"
OUTPUT = ROOT / "BAO_CAO_DO_AN_CUOI_KY_HEALWAY_HOAN_THIEN_V3.docx"
FIGURE_IMAGES = {
    1: (ROOT / "report-images" / "01-start.png", "Hình 1. Màn hình đăng nhập HealWay trên Android Emulator", 6.3),
    2: (ROOT.parent / "build" / "report-preview" / "patient-profiles.png", "Hình 2. Danh sách hồ sơ người bệnh (dữ liệu DEMO)", 6.3),
    3: (ROOT.parent / "build" / "report-preview" / "doctor-detail.png", "Hình 3. Thông tin bác sĩ theo chuyên khoa (dữ liệu DEMO)", 6.3),
    4: (ROOT.parent / "build" / "report-preview" / "appointment-time.png", "Hình 4. Chọn bác sĩ, ngày và khung giờ khám (dữ liệu DEMO)", 6.3),
    5: (ROOT.parent / "build" / "report-preview" / "appointment-history.png", "Hình 5. Lịch sử lịch hẹn của người bệnh (dữ liệu DEMO)", 6.3),
    6: (
        (ROOT.parent / "build" / "report-preview" / "payment-history.png", "Hình 6a. Lịch sử thanh toán trong giao diện HealWay (dữ liệu DEMO)", 6.3),
        (ROOT.parent / "build" / "report-preview" / "payos-checkout.png", "Hình 6b. Tạo phiên PayOS, hiển thị mã QR và kiểm tra trạng thái (giao diện DEMO)", 6.3),
    ),
    7: (
        (ROOT.parent / "build" / "report-preview" / "feedback.png", "Hình 7a. Màn hình phản hồi ý kiến khách hàng (giao diện DEMO)", 6.3),
        (ROOT.parent / "build" / "report-preview" / "chatbot.png", "Hình 7b. Hội thoại với Trợ lý HealWay (giao diện DEMO)", 6.3),
    ),
    8: (ROOT.parent / "build" / "report-preview" / "clinical-profiles.png", "Hình 8. Danh sách hồ sơ xem kết quả cận lâm sàng (dữ liệu DEMO)", 6.3),
    9: (ROOT.parent / ".dart_tool" / "doctor_slot_details_top.png", "Hình 9. Chi tiết ca khám trong không gian bác sĩ (dữ liệu DEMO)", 6.0),
    10: (ROOT.parent / ".dart_tool" / "doctor_reports_preview_top.png", "Hình 10. Màn hình báo cáo trong không gian bác sĩ (dữ liệu DEMO)", 6.0),
    11: (ROOT.parent / "build" / "admin-preview" / "overview-390.png", "Hình 11. Dashboard quản trị HealWay (dữ liệu DEMO)", 6.3),
    12: (ROOT.parent / "build" / "admin-preview" / "section-5-390.png", "Hình 12. Danh sách bác sĩ trên giao diện điện thoại (DEMO)", 6.3),
    13: (ROOT.parent / "build" / "admin-preview" / "section-2-390.png", "Hình 13. Danh sách lịch hẹn trên giao diện điện thoại (DEMO)", 6.3),
    14: (ROOT.parent / ".dart_tool" / "doctor_dashboard_preview.png", "Hình 14. Dashboard của vai trò bác sĩ (dữ liệu DEMO)", 6.0),
}

BLUE = "17365D"
PALE_BLUE = "EAF2F8"
LIGHT_GRAY = "F5F5F5"
MID_GRAY = "D9E2F3"
BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(89, 89, 89)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    p_pr.append(keep)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Trang ")
    run.font.name = "Times New Roman"
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def set_run_font(run, size=None, bold=None, italic=None, color=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = color


def add_inline(paragraph, text, size=13, color=BLACK, bold=False, italic=False):
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=size, bold=True, color=color)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size=size, bold=bold, italic=italic, color=color)
            run.font.name = "Courier New"
            run._element.rPr.rFonts.set(qn("w:ascii"), "Courier New")
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=size, color=color)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.35)
    section.bottom_margin = Cm(2.1)
    section.left_margin = Cm(2.7)
    section.right_margin = Cm(2.2)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(13)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(6)

    for name, size, before, after in (("Title", 18, 0, 18), ("Heading 1", 16, 16, 10), ("Heading 2", 14, 12, 7), ("Heading 3", 13, 10, 5)):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    footer = section.footer
    add_page_number(footer.paragraphs[0])


def add_cover(doc):
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(p, "BỘ GIÁO DỤC VÀ ĐÀO TẠO", size=14)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(p, "[TÊN TRƯỜNG]", size=14, color=GRAY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(p, "[TÊN KHOA/VIỆN]", size=14, color=GRAY)

    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(p, "BÁO CÁO ĐỒ ÁN CUỐI KỲ", size=18)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(p, "HỌC PHẦN: LẬP TRÌNH TRÊN THIẾT BỊ DI ĐỘNG (ECMP177)", size=14)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(22)
    p.paragraph_format.space_after = Pt(28)
    add_inline(p, "ỨNG DỤNG ĐẶT LỊCH KHÁM BỆNH HEALWAY", size=18, bold=True)

    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    table.columns[0].width = Cm(8.2)
    table.columns[1].width = Cm(3.0)
    table.columns[2].width = Cm(3.0)
    headers = ["Thành viên thực hiện", "MSSV", "Lớp"]
    for i, value in enumerate(headers):
        cell = table.cell(0, i)
        set_cell_shading(cell, BLUE)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(p, value, size=10, color=RGBColor(255, 255, 255))
    for name in ("[Họ tên thành viên 1]", "[Họ tên thành viên 2]", "[Bổ sung nếu có]"):
        cells = table.add_row().cells
        values = [name, "[MSSV]", "[Lớp]"]
        for i, value in enumerate(values):
            set_cell_border(cells[i])
            set_cell_margins(cells[i])
            set_cell_shading(cells[i], "FFFFFF" if len(table.rows) % 2 else PALE_BLUE)
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i else WD_ALIGN_PARAGRAPH.LEFT
            add_inline(p, value, size=10, color=GRAY)

    for _ in range(3):
        doc.add_paragraph()
    for label, value in (("Giảng viên phụ trách", "[Họ tên giảng viên]"), ("Nhóm", "[Số nhóm]"), ("Học kỳ", "[Học kỳ] - Năm học [20xx-20xx]")):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(4.0)
        add_inline(p, f"{label}: ", size=13, bold=True)
        add_inline(p, value, size=13, color=GRAY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(30)
    add_inline(p, "TP. Hồ Chí Minh, tháng [..] năm [20..]", size=13, italic=True, color=GRAY)
    doc.add_page_break()


def parse_table(lines, index):
    rows = []
    while index < len(lines) and lines[index].startswith("|"):
        cells = [part.strip() for part in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            rows.append(cells)
        index += 1
    return rows, index


def add_table(doc, rows):
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = doc.add_table(rows=1, cols=columns)
    table.autofit = True
    table.style = "Table Grid"
    for row_index, row in enumerate(rows):
        target = table.rows[0] if row_index == 0 else table.add_row()
        for column in range(columns):
            text = row[column] if column < len(row) else ""
            cell = target.cells[column]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            if row_index == 0:
                set_cell_shading(cell, BLUE)
            elif row_index % 2 == 0:
                set_cell_shading(cell, PALE_BLUE)
            else:
                set_cell_shading(cell, "FFFFFF")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.12
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if column == 0 and columns >= 3 else WD_ALIGN_PARAGRAPH.LEFT
            add_inline(p, text, size=8.5 if columns >= 4 else 9)
            if row_index == 0:
                for run in p.runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.bold = True
        if row_index == 0:
            set_repeat_table_header(target)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_placeholder_figure(doc, text):
    match = re.search(r"Hình\s+(\d+)", text)
    figures = FIGURE_IMAGES.get(int(match.group(1))) if match else None
    if figures is None:
        return
    if isinstance(figures[0], Path):
        figures = (figures,)
    for path, caption, width in figures:
        if not path.is_file():
            raise FileNotFoundError(f"Thiếu ảnh báo cáo: {path}")
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        p.add_run().add_picture(str(path), width=Cm(width))
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(8)
        add_inline(p, caption, size=10.5, italic=True, color=GRAY)


def build_body(doc, markdown):
    start = markdown.index("## 2. Kết quả thực nghiệm")
    markdown = markdown[start:]
    lines = markdown.splitlines()
    i = 0
    skip_instruction_paragraphs = (
        "Khi dán sang Word,",
        "Có thể chèn một đoạn ngắn",
        "Nếu nhóm thực sự đã chạy",
        "Không đưa vào báo cáo",
    )
    stop_at = "## Phụ lục"
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if stripped == stop_at:
            break
        if not stripped or re.fullmatch(r"-{3,}", stripped):
            i += 1
            continue
        if stripped.startswith("|"):
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            continue
        if stripped.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            add_inline(p, stripped[4:], size=14, bold=True)
            set_keep_with_next(p)
        elif stripped.startswith("#### "):
            p = doc.add_paragraph(style="Heading 3")
            add_inline(p, stripped[5:], size=13, bold=True)
            set_keep_with_next(p)
        elif stripped.startswith("## "):
            p = doc.add_paragraph(style="Heading 1")
            add_inline(p, stripped[3:], size=16, bold=True)
            set_keep_with_next(p)
        elif stripped.startswith("> "):
            content = stripped[2:]
            if not any(content.startswith(prefix) for prefix in skip_instruction_paragraphs):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.75)
                p.paragraph_format.right_indent = Cm(0.45)
                p.paragraph_format.space_before = Pt(5)
                p.paragraph_format.space_after = Pt(7)
                add_inline(p, content, size=12, italic=True, color=GRAY)
        elif stripped.startswith("**[Chèn Hình"):
            add_placeholder_figure(doc, stripped)
        elif re.match(r"^\d+\. ", stripped):
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, re.sub(r"^\d+\. ", "", stripped), size=12)
        elif any(stripped.startswith(prefix) for prefix in skip_instruction_paragraphs):
            pass
        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent = Cm(0.75)
            p.paragraph_format.space_after = Pt(6)
            add_inline(p, stripped, size=13)
        i += 1


def main():
    markdown = SOURCE.read_text(encoding="utf-8")
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    build_body(doc, markdown)
    doc.core_properties.title = "Báo cáo đồ án cuối kỳ HealWay"
    doc.core_properties.subject = "Ứng dụng đặt lịch khám bệnh HealWay"
    doc.core_properties.author = "[Nhóm thực hiện]"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
